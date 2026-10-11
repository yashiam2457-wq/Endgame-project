"""
Enterprise GraphRAG — Ingestion Package.

Consolidated Ingestion Pipeline: DocumentLoader, TextExtractor, DocumentCleaner, MetadataExtractor, IngestionPipeline.
"""

from __future__ import annotations

import asyncio
import json
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator, Optional
from uuid import uuid4

from core import (
    AppConfig,
    DocumentParsingError,
    DuplicateDocumentError,
    UnsupportedFileTypeError,
    app_logger,
)
from models import Document, DocumentMetadata, DocumentStatus, SupportedFileType
from utils import SUPPORTED_EXTENSIONS, clean_document_text, compute_sha256, count_tokens_approx, iter_supported_files, timer

# ── SECTION 1: Document Loader ───────────────────────────────────────────────

_EXT_TO_TYPE: dict[str, SupportedFileType] = {
    ".pdf": SupportedFileType.PDF,
    ".docx": SupportedFileType.DOCX,
    ".doc": SupportedFileType.DOCX,
    ".txt": SupportedFileType.TXT,
    ".md": SupportedFileType.MARKDOWN,
    ".markdown": SupportedFileType.MARKDOWN,
    ".html": SupportedFileType.HTML,
    ".htm": SupportedFileType.HTML,
}


class DocumentLoader:
    """Discovers and creates Document objects from file paths."""

    def load_file(self, file_path: Path) -> Document:
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = file_path.suffix.lower()
        if ext not in _EXT_TO_TYPE:
            raise UnsupportedFileTypeError(
                f"Unsupported file type: {ext}",
                details={"supported": list(SUPPORTED_EXTENSIONS)},
            )

        sha256 = compute_sha256(file_path)
        file_type = _EXT_TO_TYPE[ext]

        doc = Document(
            sha256=sha256,
            filename=file_path.name,
            file_type=file_type,
            file_path=file_path,
        )
        doc.metadata.file_size_bytes = file_path.stat().st_size
        doc.metadata.file_type = ext
        doc.metadata.source_path = str(file_path)

        app_logger.debug(f"Loaded document: {file_path.name} ({sha256[:8]}...)")
        return doc

    def load_directory(
        self, directory: Path, recursive: bool = True
    ) -> Iterator[Document]:
        if not directory.is_dir():
            raise NotADirectoryError(f"Not a directory: {directory}")

        found = 0
        for file_path in iter_supported_files(directory, recursive=recursive):
            try:
                yield self.load_file(file_path)
                found += 1
            except (UnsupportedFileTypeError, OSError) as exc:
                app_logger.warning(f"Skipping {file_path}: {exc}")

        app_logger.info(f"Directory scan complete: {found} documents found in {directory}")

    def load(self, path: Path, recursive: bool = True) -> Iterator[Document]:
        if path.is_file():
            yield self.load_file(path)
        elif path.is_dir():
            yield from self.load_directory(path, recursive=recursive)
        else:
            raise FileNotFoundError(f"Path does not exist: {path}")


# ── SECTION 2: Text Extractor ────────────────────────────────────────────────


class TextExtractor:
    """Dispatches text extraction to the appropriate parser."""

    def __init__(self, llm: Optional[Any] = None) -> None:
        self._llm = llm

    def extract(self, document: Document) -> Document:
        if document.file_path is None:
            raise DocumentParsingError("Document has no file_path set.")

        extractor_map = {
            SupportedFileType.PDF: self._extract_pdf,
            SupportedFileType.DOCX: self._extract_docx,
            SupportedFileType.TXT: self._extract_txt,
            SupportedFileType.MARKDOWN: self._extract_markdown,
            SupportedFileType.HTML: self._extract_html,
        }

        extractor = extractor_map.get(document.file_type)
        if extractor is None:
            raise DocumentParsingError(
                f"No extractor for file type: {document.file_type}"
            )

        try:
            text, metadata = extractor(document.file_path)
            document.raw_text = text
            document.metadata = metadata
            app_logger.debug(
                f"Extracted {len(text)} chars from {document.filename}"
            )
            return document
        except DocumentParsingError:
            raise
        except Exception as exc:
            raise DocumentParsingError(
                f"Extraction failed for {document.filename}: {exc}",
                cause=exc,
            ) from exc

    def _extract_pdf(self, path: Path) -> tuple[str, DocumentMetadata]:
        # 1. Priority check: If a matching .txt sidecar exists, load text from it first
        txt_sidecar1 = path.with_suffix(".txt")
        txt_sidecar2 = path.parent / f"{path.name}.txt"
        for sidecar in (txt_sidecar1, txt_sidecar2):
            if sidecar.exists() and sidecar.is_file() and sidecar.stat().st_size > 0:
                try:
                    txt_content = sidecar.read_text(encoding="utf-8", errors="ignore").strip()
                    if txt_content:
                        app_logger.info(f"Loaded text for PDF from txt sidecar: {sidecar.name}")
                        meta = DocumentMetadata(
                            title=path.stem,
                            file_size_bytes=path.stat().st_size,
                            file_type="pdf",
                            source_path=str(path),
                            word_count=len(txt_content.split()),
                        )
                        return txt_content, meta
                except Exception as exc:
                    app_logger.debug(f"Failed reading sidecar {sidecar.name}: {exc}")

        # 2. Use Enterprise Offline OCR Subsystem
        try:
            from ocr_pipeline import OCRProcessingPipeline
            pipeline = OCRProcessingPipeline()
            ocr_doc = pipeline.process_pdf(path)
            extracted_text = ocr_doc.reconstructed_markdown or ocr_doc.reconstructed_text
            if extracted_text and len(extracted_text.strip()) > 10:
                meta = DocumentMetadata(
                    title=path.stem,
                    page_count=ocr_doc.total_pages,
                    file_size_bytes=path.stat().st_size,
                    file_type="pdf",
                    source_path=str(path),
                    word_count=len(extracted_text.split()),
                    char_count=len(extracted_text),
                )
                app_logger.info(
                    f"Successfully processed PDF via OCR Pipeline: {path.name} "
                    f"({ocr_doc.doc_type.value.upper()}, {ocr_doc.total_pages} pages, {meta.word_count} words)"
                )
                return extracted_text, meta
        except Exception as ocr_pipeline_exc:
            app_logger.warning(
                f"OCR Pipeline encounter fallback warning on {path.name}: {ocr_pipeline_exc}. Proceeding with standard PyMuPDF parser."
            )

        # 3. Extract plain text directly using PyMuPDF (fitz) fallback

        try:
            import fitz
        except ImportError as exc:
            raise DocumentParsingError(
                "PyMuPDF not installed. Run: pip install PyMuPDF", cause=exc
            ) from exc

        # Auto-configure Tesseract binary path on Windows if needed
        try:
            import pytesseract
            import shutil
            if not shutil.which("tesseract"):
                possible_paths = [
                    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                    str(Path.home() / r"AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
                    str(Path.home() / r"AppData\Local\Tesseract-OCR\tesseract.exe"),
                ]
                for p in possible_paths:
                    if os.path.exists(p):
                        pytesseract.pytesseract.tesseract_cmd = p
                        app_logger.info(f"Configured Tesseract binary path: {p}")
                        break
        except Exception:
            pass

        try:
            doc = fitz.open(str(path))
            metadata = DocumentMetadata(
                title=doc.metadata.get("title") or path.stem,
                author=doc.metadata.get("author") or None,
                page_count=doc.page_count,
                file_size_bytes=path.stat().st_size,
                file_type="pdf",
                source_path=str(path),
            )

            text_parts: list[str] = []
            for page_idx, page in enumerate(doc):
                page_text = ""
                # Method A: Direct plain text extraction
                try:
                    page_text = page.get_text("text").strip()
                except Exception:
                    pass

                # Method B: Recover text blocks or words if get_text("text") returned sparse/empty
                if len(page_text) < 30:
                    try:
                        blocks = page.get_text("blocks")
                        block_texts = [b[4].strip() for b in blocks if len(b) >= 5 and b[4].strip()]
                        if block_texts:
                            page_text = "\n".join(block_texts).strip()
                    except Exception:
                        pass

                # Method C: Attempt Tesseract OCR if text is still sparse
                if len(page_text) < 30:
                    try:
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        try:
                            import io
                            from PIL import Image
                            import pytesseract
                            img = Image.open(io.BytesIO(img_bytes))
                            ocr_text = pytesseract.image_to_string(img).strip()
                            if ocr_text:
                                page_text = f"--- Page {page_idx + 1} (OCR/Diagram) ---\n{ocr_text}"
                        except Exception as ocr_exc:
                            app_logger.debug(
                                f"Image OCR skipped on page {page_idx + 1} of {path.name}: {ocr_exc}"
                            )
                    except Exception:
                        pass

                # Method D: Diagram & Image Vector Structure Inspection if no text extracted
                if len(page_text) < 30:
                    try:
                        images = page.get_images()
                        drawings = page.get_drawings()
                        diag_labels = []
                        for d in drawings:
                            for item in d.get("items", []):
                                if len(item) > 1 and isinstance(item[1], str):
                                    diag_labels.append(item[1])
                        img_desc = f"[Page {page_idx + 1}: Image-Based Content with {len(images)} embedded images and {len(drawings)} diagram visual paths]"
                        if diag_labels:
                            img_desc += f"\nDiagram Labels: {', '.join(diag_labels[:20])}"
                        page_text = f"--- Page {page_idx + 1} (Image/Diagram) ---\n{img_desc}"
                    except Exception:
                        pass

                # Method E: Gemma 4 Vision — send page pixmap to Ollama for intelligent description
                # This is the final fallback for image-only/diagram pages
                if len(page_text) < 100 and self._llm is not None:
                    try:
                        pix = page.get_pixmap(dpi=120)
                        img_bytes = pix.tobytes("png")
                        app_logger.info(
                            f"Sending page {page_idx + 1} of {path.name} to Gemma 4 for vision parsing..."
                        )
                        gemma_desc = self._llm.describe_diagram(img_bytes)
                        if gemma_desc and len(gemma_desc) > 20:
                            prefix = f"--- Page {page_idx + 1} (Gemma 4 Vision Analysis) ---\n"
                            page_text = prefix + gemma_desc
                            app_logger.info(
                                f"Gemma 4 vision parsed page {page_idx + 1} of {path.name}: "
                                f"{len(gemma_desc)} chars extracted."
                            )
                    except Exception as gemma_exc:
                        app_logger.debug(
                            f"Gemma 4 vision skipped on page {page_idx + 1} of {path.name}: {gemma_exc}"
                        )

                if page_text:
                    text_parts.append(f"--- Page {page_idx + 1} ---\n{page_text}")

            full_text = "\n\n".join(text_parts).strip()
            doc.close()
        except Exception as pdf_exc:
            app_logger.warning(f"PDF extraction failed for {path.name}: {pdf_exc}")
            full_text = ""
            metadata = DocumentMetadata(
                title=path.stem,
                file_size_bytes=path.stat().st_size,
                file_type="pdf",
                source_path=str(path),
            )

        # 4. Fallback if PDF contained no readable text
        if not full_text:
            full_text = f"Document: {path.name}\n(Image/Diagram PDF document: {path.name})"

        metadata.word_count = len(full_text.split())
        return full_text, metadata

    def _extract_docx(self, path: Path) -> tuple[str, DocumentMetadata]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise DocumentParsingError(
                "python-docx not installed. Run: pip install python-docx", cause=exc
            ) from exc

        docx = DocxDocument(str(path))
        parts: list[str] = []

        for para in docx.paragraphs:
            if para.text.strip():
                parts.append(para.text)

        for table in docx.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)

        full_text = "\n".join(parts)
        core_props = docx.core_properties
        metadata = DocumentMetadata(
            title=core_props.title or None,
            author=core_props.author or None,
            created_at=core_props.created,
            modified_at=core_props.modified,
            word_count=len(full_text.split()),
            file_size_bytes=path.stat().st_size,
            file_type="docx",
            source_path=str(path),
        )
        return full_text, metadata

    def _extract_txt(self, path: Path) -> tuple[str, DocumentMetadata]:
        import chardet

        raw_bytes = path.read_bytes()
        detected = chardet.detect(raw_bytes)
        encoding = detected.get("encoding") or "utf-8"

        try:
            text = raw_bytes.decode(encoding, errors="replace")
        except Exception:
            text = raw_bytes.decode("utf-8", errors="replace")

        metadata = DocumentMetadata(
            word_count=len(text.split()),
            file_size_bytes=path.stat().st_size,
            file_type="txt",
            source_path=str(path),
        )
        return text, metadata

    def _extract_markdown(self, path: Path) -> tuple[str, DocumentMetadata]:
        import re

        text = path.read_text(encoding="utf-8", errors="replace")
        text = re.sub(r"```[\s\S]*?```", "", text)
        text = re.sub(r"`[^`]+`", "", text)
        text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
        text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
        text = re.sub(r"!\[([^\]]*)\]\([^\)]+\)", r"\1", text)
        text = re.sub(r"[*_]{1,3}([^*_]+)[*_]{1,3}", r"\1", text)

        metadata = DocumentMetadata(
            word_count=len(text.split()),
            file_size_bytes=path.stat().st_size,
            file_type="md",
            source_path=str(path),
        )
        return text, metadata

    def _extract_html(self, path: Path) -> tuple[str, DocumentMetadata]:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise DocumentParsingError(
                "beautifulsoup4 not installed. Run: pip install beautifulsoup4",
                cause=exc,
            ) from exc

        html_content = path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(html_content, "lxml")

        for tag in soup(["script", "style", "meta", "head"]):
            tag.decompose()

        title = soup.title.string if soup.title else None
        text = soup.get_text(separator="\n", strip=True)

        metadata = DocumentMetadata(
            title=title,
            word_count=len(text.split()),
            file_size_bytes=path.stat().st_size,
            file_type="html",
            source_path=str(path),
        )
        return text, metadata


# ── SECTION 3: Document Cleaner ─────────────────────────────────────────────


class DocumentCleaner:
    """Applies cleaning pipeline to document raw text."""

    def clean(self, document: Document) -> Document:
        document.cleaned_text = clean_document_text(document.raw_text)
        return document


# ── SECTION 4: Metadata Extractor ────────────────────────────────────────────


class MetadataExtractor:
    """Enriches document metadata after text extraction."""

    def enrich(self, document: Document) -> Document:
        text = document.cleaned_text or document.raw_text
        if not document.metadata.word_count:
            document.metadata.word_count = len(text.split())

        document.metadata.extra["token_count_approx"] = count_tokens_approx(text)
        document.metadata.extra["char_count"] = len(text)

        if not document.metadata.title:
            first_line = next(
                (line.strip() for line in text.splitlines() if line.strip()), None
            )
            if first_line and len(first_line) < 200:
                document.metadata.title = first_line

        return document


# ── SECTION 5: Ingestion Pipeline ───────────────────────────────────────────


class IngestionPipeline:
    """Full offline indexing pipeline orchestrator."""

    def __init__(
        self,
        settings: AppConfig,
        sqlite: Any,
        faiss_store: Any,
        neo4j_client: Any,
        encoder: Any,
        llm: Any,
        cache: Any,
    ) -> None:
        self._settings = settings
        self._sqlite = sqlite
        self._faiss = faiss_store
        self._neo4j = neo4j_client
        self._encoder = encoder
        self._llm = llm
        self._cache = cache

        self._loader = DocumentLoader()
        self._extractor = TextExtractor(llm=self._llm)
        self._cleaner = DocumentCleaner()
        self._meta_extractor = MetadataExtractor()
        self._job_progress: dict[str, dict[str, Any]] = {}

        self._executor = ThreadPoolExecutor(
            max_workers=settings.performance.worker_threads
            if hasattr(settings, "performance")
            else 4
        )

    def update_job_progress(
        self,
        job_id: str,
        stage: str,
        percent: int,
        detail: str = "",
        filename: str = "",
        status: str = "processing",
    ) -> None:
        prev = self._job_progress.get(job_id, {})
        fname = filename or prev.get("filename", "")
        self._job_progress[job_id] = {
            "job_id": job_id,
            "filename": fname,
            "stage": stage,
            "progress_percent": percent,
            "detail": detail,
            "status": status,
            "updated_at": datetime.utcnow().isoformat(),
        }

    def get_job_progress(self, job_id: str) -> dict[str, Any]:
        if job_id in self._job_progress:
            return self._job_progress[job_id]
        return {
            "job_id": job_id,
            "filename": "Unknown",
            "stage": "Initializing",
            "progress_percent": 0,
            "detail": "Job queued",
            "status": "pending",
        }

    async def ingest_file(
        self,
        file_path: Path,
        force: bool = False,
        job_id: Optional[str] = None,
    ) -> dict[str, Any]:
        job_id = job_id or str(uuid4())
        app_logger.info(f"[{job_id}] Starting ingestion: {file_path.name}")
        self.update_job_progress(
            job_id, "Loading Document", 5, "Reading file from storage...", filename=file_path.name
        )

        with timer(f"ingestion:{file_path.name}") as t:
            doc = self._loader.load_file(file_path)

            self.update_job_progress(job_id, "Deduplication Check", 10, "Checking fingerprint...")
            await asyncio.sleep(0.01)
            existing = self._sqlite.get_document_by_sha256(doc.sha256)
            if existing:
                doc.id = existing["id"]
                ex_status = existing.get("status", "")
                if ex_status == DocumentStatus.INDEXED.value and not force:
                    self.update_job_progress(
                        job_id, "Skipped", 100, "Document already indexed", status="completed"
                    )
                    await asyncio.sleep(0.01)
                    raise DuplicateDocumentError(
                        f"Document already indexed: {file_path.name}",
                        details={"existing_id": existing["id"], "sha256": doc.sha256},
                    )
                else:
                    app_logger.info(
                        f"Re-ingesting document {file_path.name} (id: {existing['id']}, status: {ex_status})"
                    )

            self.update_job_progress(
                job_id, "Text & Diagram Extraction", 20, "Extracting text, OCR & Gemma diagrams..."
            )
            await asyncio.sleep(0.01)
            doc = self._extractor.extract(doc)

            self.update_job_progress(job_id, "Cleaning Text", 30, "Normalizing document whitespace...")
            await asyncio.sleep(0.01)
            doc = self._cleaner.clean(doc)

            self.update_job_progress(job_id, "Metadata Enrichment", 35, "Extracting document properties...")
            await asyncio.sleep(0.01)
            doc = self._meta_extractor.enrich(doc)

            doc.status = DocumentStatus.PROCESSING
            self._persist_document(doc)

            self.update_job_progress(job_id, "Text Chunking", 45, "Splitting document into text chunks...")
            chunks = await self._chunk_document(doc)
            for c in chunks:
                c.document_filename = doc.filename
                if isinstance(c.metadata, dict):
                    c.metadata["document_filename"] = doc.filename
            app_logger.info(f"[{job_id}] {len(chunks)} chunks created.")

            self.update_job_progress(
                job_id, "Generating Embeddings", 60, f"Encoding {len(chunks)} chunks into vector space..."
            )
            await asyncio.sleep(0.01)
            embeddings = await self._generate_embeddings(chunks, job_id)

            self.update_job_progress(job_id, "FAISS Index Insertion", 70, "Storing vectors in FAISS index...")
            await asyncio.sleep(0.01)
            self._faiss.add_chunks(chunks, embeddings)

            self.update_job_progress(
                job_id, "Gemma Entity Extraction", 70, f"Extracting Knowledge Graph from {len(chunks)} chunks..."
            )
            await asyncio.sleep(0.01)
            extraction_results = await self._extract_entities(chunks, doc, job_id)

            all_entities = []
            all_relationships = []
            for res in extraction_results:
                all_entities.extend(res.entities)
                all_relationships.extend(res.relationships)

            self.update_job_progress(
                job_id, "Knowledge Graph Construction", 90, f"Connecting {len(all_entities)} entities in Neo4j..."
            )
            entity_count, rel_count = await self._build_graph(extraction_results, doc)

            from graph import CommunityDetector
            detector = CommunityDetector()
            communities = detector.detect_communities(all_entities, all_relationships)

            from generation import MultiLevelSummarizer
            summarizer = MultiLevelSummarizer(self._llm)
            summary_package = summarizer.generate_all_summaries(
                doc.id, all_entities, all_relationships, communities
            )

            self.update_job_progress(
                job_id, "Multi-Vector Indexing", 95, "Storing multi-vector entity/community embeddings..."
            )
            from embedding import MultiVectorEmbedder
            mv_embedder = MultiVectorEmbedder(self._encoder)

            entity_vecs = mv_embedder.embed_entities(all_entities)
            rel_vecs = mv_embedder.embed_relationships(all_relationships)
            comm_vecs = mv_embedder.embed_communities(communities)

            if all_entities and entity_vecs:
                self._faiss.add_vectors(
                    [e.id for e in all_entities],
                    entity_vecs,
                    namespace="entity",
                    metadata_list=[{"name": e.name, "type": e.entity_type.value, "document_id": doc.id} for e in all_entities],
                )

            if all_relationships and rel_vecs:
                self._faiss.add_vectors(
                    [r.id for r in all_relationships],
                    rel_vecs,
                    namespace="relationship",
                    metadata_list=[{"source": r.source_entity_name, "target": r.target_entity_name, "type": r.relationship_type.value, "document_id": doc.id} for r in all_relationships],
                )

            if communities and comm_vecs:
                self._faiss.add_vectors(
                    [c.id for c in communities],
                    comm_vecs,
                    namespace="community",
                    metadata_list=[{"title": c.title, "document_id": doc.id} for c in communities],
                )

            from graph import GraphBuilder
            builder = GraphBuilder(self._neo4j)
            builder.build_communities(communities)

            doc.status = DocumentStatus.INDEXED
            self._sqlite.update_document_status(doc.id, DocumentStatus.INDEXED)

            self.update_job_progress(
                job_id, "Completed", 100, f"Successfully indexed {doc.filename}", status="completed"
            )

        app_logger.info(
            f"[{job_id}] Ingestion complete: {doc.filename} — "
            f"{len(chunks)} chunks, {entity_count} entities, "
            f"{rel_count} relationships, {len(communities)} communities in {t['elapsed_ms']:.0f} ms"
        )

        return {
            "job_id": job_id,
            "document_id": doc.id,
            "filename": doc.filename,
            "chunk_count": len(chunks),
            "entity_count": entity_count,
            "relationship_count": rel_count,
            "community_count": len(communities),
            "summary_package": summary_package.model_dump(mode="json"),
            "elapsed_ms": t["elapsed_ms"],
            "status": "completed",
        }

    def _persist_document(self, doc: Document) -> None:
        now = datetime.utcnow().isoformat()
        self._sqlite.upsert_document(
            {
                "id": doc.id,
                "sha256": doc.sha256,
                "filename": doc.filename,
                "file_type": doc.file_type.value,
                "file_path": str(doc.file_path) if doc.file_path else None,
                "status": doc.status.value,
                "title": doc.metadata.title,
                "author": doc.metadata.author,
                "page_count": doc.metadata.page_count,
                "word_count": doc.metadata.word_count,
                "file_size_bytes": doc.metadata.file_size_bytes,
                "version": doc.version,
                "metadata_json": json.dumps(doc.metadata.extra),
                "created_at": now,
                "updated_at": now,
            }
        )

    async def _chunk_document(self, doc: Document) -> list[Any]:
        from chunking import ChunkerFactory

        chunker = ChunkerFactory.create(self._settings.chunking)
        chunks = chunker.chunk(doc)

        if not chunks:
            fallback_text = doc.text.strip() if doc.text and doc.text.strip() else f"Document content for {doc.filename}"
            from models import Chunk, ChunkType
            c = Chunk(
                document_id=doc.id,
                sequence_index=0,
                text=fallback_text,
                chunk_type=ChunkType.PARAGRAPH,
                token_count=count_tokens_approx(fallback_text),
                section_header=doc.filename,
            )
            chunks = [c]

        chunk_records = [
            {
                "id": chunk.id,
                "document_id": chunk.document_id,
                "sequence_index": chunk.sequence_index,
                "text": chunk.text,
                "chunk_type": chunk.chunk_type.value,
                "parent_chunk_id": chunk.parent_chunk_id,
                "start_char": chunk.start_char,
                "end_char": chunk.end_char,
                "page_number": chunk.page_number,
                "section_header": chunk.section_header,
                "token_count": chunk.token_count,
                "metadata_json": json.dumps(chunk.metadata),
                "created_at": datetime.utcnow().isoformat(),
            }
            for chunk in chunks
        ]
        self._sqlite.insert_chunks(chunk_records)
        return chunks

    async def _generate_embeddings(
        self, chunks: list[Any], job_id: str
    ) -> list[list[float]]:
        from embedding import BatchEmbedder

        app_logger.info(f"[{job_id}] Generating embeddings for {len(chunks)} chunks...")
        embedder = BatchEmbedder(self._encoder, self._cache)
        texts = [chunk.text for chunk in chunks]
        return embedder.embed_batch(texts, cache_prefix="chunk")

    async def _extract_entities(
        self, chunks: list[Any], doc: Document, job_id: str
    ) -> list[Any]:
        from generation import EntityExtractor

        app_logger.info(f"[{job_id}] Extracting entities from {len(chunks)} chunks...")
        extractor = EntityExtractor(self._llm, self._cache)
        results = []
        total_chunks = len(chunks)
        for idx, chunk in enumerate(chunks, start=1):
            try:
                result = extractor.extract(chunk, doc.id)
                results.append(result)
            except Exception as exc:
                app_logger.warning(
                    f"[{job_id}] Entity extraction failed for chunk {chunk.id}: {exc}"
                )
            if total_chunks > 0:
                pct = 70 + int((idx / total_chunks) * 18)
                self.update_job_progress(
                    job_id, "Gemma 4 Entity Extraction", min(88, pct), f"Extracting Knowledge Graph from chunk {idx}/{total_chunks}..."
                )
            await asyncio.sleep(0.01)
        return results

    async def _build_graph(
        self, extraction_results: list[Any], doc: Document
    ) -> tuple[int, int]:
        from graph import GraphBuilder

        builder = GraphBuilder(self._neo4j)
        entity_count, rel_count = builder.build_from_extractions(
            extraction_results, doc
        )
        return entity_count, rel_count

    async def ingest_directory(
        self, directory: Path, force: bool = False
    ) -> list[dict[str, Any]]:
        try:
            from tqdm import tqdm
            has_tqdm = True
        except ImportError:
            has_tqdm = False

        docs_to_process = list(self._loader.load_directory(directory))
        results = []

        iterator = tqdm(docs_to_process, desc="Indexing Documents", unit="file") if has_tqdm else docs_to_process

        for doc_item in iterator:
            try:
                if has_tqdm and hasattr(iterator, "set_postfix"):
                    iterator.set_postfix(file=doc_item.filename[:20])

                result = await self.ingest_file(
                    doc_item.file_path, force=force
                )
                results.append(result)
            except DuplicateDocumentError as exc:
                app_logger.info(f"Skipping duplicate: {exc.message}")
                results.append({"filename": doc_item.filename, "status": "skipped_duplicate"})
            except Exception as exc:
                app_logger.error(f"Ingestion failed for {doc_item.filename}: {exc}")
                results.append({"filename": doc_item.filename, "status": "failed", "error": str(exc)})
        return results


# ── SECTION 6: Backward Compatibility Submodule Aliases ─────────────────────

for _mod in ["cleaner", "extractor", "loader", "metadata", "pipeline"]:
    sys.modules[f"ingestion.{_mod}"] = sys.modules[__name__]
